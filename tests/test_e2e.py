"""End-to-end June 2024 acceptance from raw inputs to workbook (§17). Uses the real ledger and
every real BBVA statement; skips if inputs are absent."""
from decimal import Decimal

from openpyxl import load_workbook

from conftest import ROOT
from recon import store
from recon.config import Config
from recon.pipeline import run_period
from recon.output.workbook import write_workbook

# §17.2 bank open/close per account — the verified acceptance figures
BANK_OPEN_CLOSE = {
    "1112-01-001-00": ("98971.26", "346100.50"),
    "1112-01-002-00": ("18826.19", "19066.95"),
    "1112-01-003-00": ("30870.48", "13622.18"),
    "1112-01-004-00": ("152626.97", "13708.18"),
    "1112-01-005-00": ("149572.20", "63032.11"),
    "1112-01-006-00": ("120289.31", "44353.04"),
    "1112-01-007-00": ("40254.91", "24963.67"),
    "1112-01-009-00": ("52061.41", "152132.36"),
    "1112-01-011-00": ("163411.52", "69214.00"),
    "1112-01-012-00": ("91601.00", "130439.60"),
    "1112-01-013-00": ("242023.88", "266310.29"),
    "1112-01-016-00": ("13000.01", "13000.01"),
    "1112-01-017-00": ("183587.20", "191453.34"),
    "1112-02-001-00": ("2167.08", "2299.05"),
}


def _run(root):
    conn = store.connect(":memory:")
    store.init_db(conn)
    config = Config(str(ROOT / "config"))
    rr = run_period(conn, config, str(root), "2024-06", root / "diag")
    return conn, config, rr


def test_june_full_entity_all_gates_green(june_data_root):
    root, n_stmts = june_data_root
    conn, config, rr = _run(root)
    assert rr.status == "ok"

    # every GATE invariant is green for the whole entity (R1 opening-variance is a report)
    gates = [r for r in rr.invariants if r.gate]
    assert all(r.passed for r in gates), [(r.id, r.ledger_account, r.detail)
                                          for r in gates if not r.passed]

    # I3/I4/I5/I6/I7 each covered all statement-backed accounts
    reconciled = [v for v in rr.views if v.stmt is not None]
    assert len(reconciled) == n_stmts == 14


def test_june_bank_balances_match_17_2(june_data_root):
    root, _ = june_data_root
    conn, config, rr = _run(root)
    got = {v.account.ledger_account: (v.stmt.opening_balance, v.stmt.closing_balance)
           for v in rr.views if v.stmt is not None}
    for acct, (o, c) in BANK_OPEN_CLOSE.items():
        assert got[acct] == (Decimal(o), Decimal(c)), acct


def test_june_closure_residual_zero_every_account(june_data_root):
    root, _ = june_data_root
    conn, config, rr = _run(root)
    for v in rr.views:
        if v.stmt is not None:
            assert v.ap.closure_residual == Decimal("0.00"), \
                f"{v.account.ledger_account}: {v.ap.closure_residual}"


def test_june_pilot_anomalies_present(june_data_root):
    # the structural findings the pilot made by hand (§12.6)
    root, _ = june_data_root
    conn, config, rr = _run(root)
    kinds = {a["kind"] for a in rr.anomalies}
    assert "currency_mismatch" in kinds          # USD account 1112-02-001-00
    assert "dormant_with_balance" in kinds       # 1112-01-016-00 at 13,000.01
    assert "sign_divergence" in kinds            # e.g. 1112-01-012-00


def test_june_workbook_opens_with_contract_sheets(june_data_root):
    root, _ = june_data_root
    conn, config, rr = _run(root)
    out = write_workbook(conn, config, rr, str(root / "wb.xlsx"))
    wb = load_workbook(out)
    assert wb.sheetnames == ["Cover", "1. Reconciliation Summary", "2. Item Validation Report",
                             "3. Investigation Register", "4. Proposed Adjusting Entries"]


def test_determinism_two_runs_identical_matches(june_data_root):
    root, _ = june_data_root

    def matches():
        _, _, rr = _run(root)
        return sorted((v.account.ledger_account, m.pass_no, str(m.ledger_amount),
                       str(m.bank_amount), tuple(w.id for w in m.ledger), tuple(w.id for w in m.bank))
                      for v in rr.views for m in v.outcome.matches)

    assert matches() == matches()
