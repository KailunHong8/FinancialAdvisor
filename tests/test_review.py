"""Exception round-trip (§14): write register -> edit -> ingest-review -> resolution attaches;
unknown id warns/skips; 'Yes' without text rejected; idempotent."""
import pytest
from openpyxl import Workbook

from recon.exceptions_io import ingest_review


def _seed_item(conn, iid):
    conn.execute("INSERT OR IGNORE INTO runs (run_id,started_at,entity,period,ledger_file,"
                 "ledger_sha256,code_version,config_sha256,status) VALUES "
                 "('R','t','E','2024-06','l','s','1.0','c','ok')")
    conn.execute(
        "INSERT OR REPLACE INTO reconciling_items (id,run_id,entity,ledger_account,period,side,"
        "direction,category,amount,source_ref,status,first_seen_period,periods_open,evidence_json)"
        " VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (iid, "R", "E", "A", "2024-06", "ledger_outstanding", "inflow", "deposit_in_transit",
         "500.00", "ledger row 1", "flagged", "2024-06", 1, "{}"))
    conn.commit()


def _register(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "3. Investigation Register"
    headers = ["item_id", "Item", "Ledger Acct", "Amount", "Found in bank?", "Found in ledger?",
               "Supporting docs needed", "Auditor conclusion", "Resolution (fill in)", "Resolved?"]
    for c, h in enumerate(headers, 1):
        ws.cell(3, c, h)
    r = 4
    for iid, resolution, resolved in rows:
        ws.cell(r, 1, iid)
        ws.cell(r, 9, resolution)
        ws.cell(r, 10, resolved)
        r += 1
    wb.save(path)
    return str(path)


def test_resolution_attaches(memdb, tmp_path):
    _seed_item(memdb, "ITEMX")
    path = _register(tmp_path / "rev.xlsx", [("ITEMX", "Cleared 02/Jul via SPEI", "Yes")])
    summary = ingest_review(memdb, path)
    assert summary["updated"] == 1
    row = memdb.execute("SELECT status, resolution, resolved_by FROM reconciling_items WHERE id='ITEMX'").fetchone()
    assert row["status"] == "resolved" and row["resolved_by"] == "user"
    assert "SPEI" in row["resolution"]


def test_unknown_id_skipped(memdb, tmp_path):
    path = _register(tmp_path / "rev.xlsx", [("NOPE", "x", "Yes")])
    summary = ingest_review(memdb, path)
    assert summary["skipped_unknown"] == 1 and summary["updated"] == 0


def test_yes_without_text_rejected(memdb, tmp_path):
    _seed_item(memdb, "ITEMY")
    path = _register(tmp_path / "rev.xlsx", [("ITEMY", "", "Yes")])
    summary = ingest_review(memdb, path)
    assert summary["rejected"] == 1 and summary["updated"] == 0


def test_idempotent(memdb, tmp_path):
    _seed_item(memdb, "ITEMZ")
    path = _register(tmp_path / "rev.xlsx", [("ITEMZ", "done", "Yes")])
    ingest_review(memdb, path)
    before = memdb.execute("SELECT resolution FROM reconciling_items WHERE id='ITEMZ'").fetchone()["resolution"]
    ingest_review(memdb, path)
    after = memdb.execute("SELECT resolution FROM reconciling_items WHERE id='ITEMZ'").fetchone()["resolution"]
    assert before == after


def test_run_reads_back_before_regenerating(memdb, config, tmp_path, monkeypatch):
    """A scheduled run must ingest hand-entered resolutions from the existing workbook BEFORE it
    re-derives, so they are not silently overwritten (§14). Uses a stubbed ingest + no ledger, so
    the run aborts right after the read-back — proving the read-back is wired ahead of it."""
    import recon.exceptions_io as eio
    from recon.pipeline import run_period, RunError

    seen = {}

    def fake_ingest(conn, path):
        seen["path"] = path
        return {"updated": 1, "skipped_unknown": 0, "rejected": 0}

    monkeypatch.setattr(eio, "ingest_review", fake_ingest)

    existing = tmp_path / "wb.xlsx"
    existing.write_bytes(b"x")  # only needs to exist; ingest is stubbed
    with pytest.raises(RunError):  # no ledger under root -> aborts, but only after read-back
        run_period(memdb, config, str(tmp_path), "2024-06", tmp_path / "diag",
                   read_back_path=str(existing))
    assert seen["path"] == str(existing)
