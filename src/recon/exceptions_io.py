"""Exception handoff read-back (§14). The engine writes Tab 3 with stable item_ids; the user
resolves items in Excel with Copilot; `recon ingest-review` reads the result back.

Resolutions are keyed by item_id (a content hash, §11.1), which is why a later `recon run`
re-derives items and re-attaches stored resolutions instead of losing them.
"""
from __future__ import annotations

import logging

import openpyxl

log = logging.getLogger("recon")

VALID_RESOLVED = {"no", "yes", "not an item"}


def ingest_review(conn, path: str) -> dict:
    """Read (item_id, Resolution, Resolved?) from a reviewed Investigation Register and update
    reconciling_items. Idempotent. Returns a summary dict."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = None
    for name in wb.sheetnames:
        if name.strip().startswith("3.") or "Investigation" in name:
            ws = wb[name]
            break
    if ws is None:
        raise ValueError("no '3. Investigation Register' sheet found in reviewed workbook")

    # locate header row and columns by label
    header_row = None
    col = {}
    for r in range(1, 8):
        labels = {str(ws.cell(r, c).value).strip().lower(): c
                  for c in range(1, ws.max_column + 1) if ws.cell(r, c).value}
        if "item_id" in labels and any("resolved" in k for k in labels):
            header_row = r
            col["item_id"] = labels["item_id"]
            col["resolution"] = next((v for k, v in labels.items() if "resolution" in k), None)
            col["resolved"] = next((v for k, v in labels.items() if "resolved?" in k
                                    or k == "resolved"), None)
            break
    if header_row is None:
        raise ValueError("could not locate item_id / Resolved? header in the register sheet")

    updated = skipped_unknown = rejected = 0
    for r in range(header_row + 1, ws.max_row + 1):
        iid = ws.cell(r, col["item_id"]).value
        if not iid:
            continue
        iid = str(iid).strip()
        resolved = str(ws.cell(r, col["resolved"]).value or "").strip().lower() if col["resolved"] else ""
        resolution = (str(ws.cell(r, col["resolution"]).value).strip()
                      if col["resolution"] and ws.cell(r, col["resolution"]).value else "")
        if resolved not in VALID_RESOLVED or resolved == "no" or resolved == "":
            continue
        exists = conn.execute("SELECT 1 FROM reconciling_items WHERE id=? LIMIT 1", (iid,)).fetchone()
        if exists is None:
            log.warning("ingest-review: unknown item_id %s — skipped", iid)
            skipped_unknown += 1
            continue
        if resolved == "yes" and not resolution:
            log.warning("ingest-review: item %s marked Yes with no Resolution — rejected", iid)
            rejected += 1
            continue
        status = "resolved" if resolved == "yes" else "resolved"  # 'Not an item' also closes it
        note = resolution if resolved == "yes" else (resolution or "Marked 'Not an item'.")
        conn.execute(
            "UPDATE reconciling_items SET resolution=?, status=?, resolved_by='user' WHERE id=?",
            (note, status, iid))
        updated += 1
    conn.commit()
    summary = {"updated": updated, "skipped_unknown": skipped_unknown, "rejected": rejected}
    log.info("ingest-review: %s", summary)
    return summary
