"""Workbook styling (§15.6). Meaning lives in columns, not in colour, so it survives CSV export
and a diff. The only colour that encodes anything is negative-amount red, applied via number
format."""
from __future__ import annotations

from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)
MONEY_FMT = "#,##0.00;[Red]-#,##0.00"
DATE_FMT = "dd/mmm/yyyy"
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws: Worksheet, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER


def money(cell) -> None:
    cell.number_format = MONEY_FMT


def date_cell(cell) -> None:
    cell.number_format = DATE_FMT


def autosize(ws: Worksheet, widths: dict[int, int]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
