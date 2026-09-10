"""CONTPAQi 'Movimientos, Auxiliares del Catálogo' parser (§9).

Grammar verified against all six 2024 sheets. Layout varies by client config, so the header row
is located by scanning for its eight labels, never hardcoded (§9.1).
"""
from __future__ import annotations

import re
from datetime import date

import openpyxl

from ..config import EntityConfig
from ..models import LedgerAccountBlock, LedgerTransaction
from ..money import q, ZERO

MONTHS_ES = {m: i for i, m in enumerate(
    ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"], start=1)}

ACCT_RE = re.compile(r"^\d{4}-\d{2}-\d{3}-\d{2}$")
DATE_RE = re.compile(r"^(\d{2})/(Ene|Feb|Mar|Abr|May|Jun|Jul|Ago|Sep|Oct|Nov|Dic)/(\d{4})$")
PERIOD_RE = re.compile(r"del\s+(\d{2})/(\w{3})/(\d{4})\s+al\s+(\d{2})/(\w{3})/(\d{4})", re.IGNORECASE)
HEADER_LABELS = ["Fecha", "Tipo", "Número", "Concepto", "Referencia", "Cargos", "Abonos", "Saldo"]
FORMAT_ASSERT = "Movimientos, Auxiliares del Catálogo"


class LedgerParseError(Exception):
    pass


def _es_date(text: str) -> date:
    m = DATE_RE.match(text.strip())
    if not m:
        raise LedgerParseError(f"bad ledger date {text!r}")
    return date(int(m.group(3)), MONTHS_ES[m.group(2)], int(m.group(1)))


def period_to_sheet(period: str) -> str:
    """'2024-06' -> '06-24'."""
    year, month = period.split("-")
    return f"{month}-{year[2:]}"


def _norm(v) -> str:
    return "" if v is None else str(v).strip()


def _find_header_row(ws, max_scan: int = 20) -> tuple[int, dict[str, int]]:
    """Return (row_index, {label: col_index}) by matching the eight header labels (§9.1)."""
    for r in range(1, max_scan + 1):
        cells = {_norm(ws.cell(r, c).value): c for c in range(1, ws.max_column + 1)}
        if all(lbl in cells for lbl in HEADER_LABELS):
            return r, {lbl: cells[lbl] for lbl in HEADER_LABELS}
    raise LedgerParseError("could not locate CONTPAQi header row (eight labels not found)")


def _cell_num(ws, r, c):
    v = ws.cell(r, c).value
    if v is None or v == "":
        return None
    if isinstance(v, str):
        v = v.replace(",", "").strip()
        if not v:
            return None
        try:
            return q(v)
        except (ArithmeticError, ValueError):
            # Excel error values ('#REF!', '#N/A') or stray labels in a numeric column -> empty
            return None
    return q(v)


class _Cols:
    def __init__(self, m: dict[str, int]):
        self.fecha = m["Fecha"]
        self.tipo = m["Tipo"]
        self.numero = m["Número"]
        self.concepto = m["Concepto"]
        self.referencia = m["Referencia"]
        self.cargos = m["Cargos"]
        self.abonos = m["Abonos"]
        self.saldo = m["Saldo"]


def parse_sheet(ws, entity: str, period: str, cfg: EntityConfig) -> list[LedgerAccountBlock]:
    header_row, colmap = _find_header_row(ws)
    cols = _Cols(colmap)

    # metadata cross-checks (§9.1)
    meta = " ".join(_norm(ws.cell(r, c).value)
                    for r in range(1, header_row) for c in range(1, ws.max_column + 1))
    if FORMAT_ASSERT not in meta:
        raise LedgerParseError(f"format assertion failed: '{FORMAT_ASSERT}' absent")
    pm = PERIOD_RE.search(meta)
    if pm:
        start = date(int(pm.group(3)), MONTHS_ES[pm.group(2).title()], int(pm.group(1)))
        got = f"{start.year:04d}-{start.month:02d}"
        if got != period:
            raise LedgerParseError(
                f"period in sheet metadata ({got}) disagrees with sheet name ({period})")

    blocks: list[LedgerAccountBlock] = []
    sheet_name = ws.title
    rollup = set(cfg.ledger_rollup_accounts)

    cur: LedgerAccountBlock | None = None
    cur_is_rollup = False
    seen_total = False   # whether the current block has passed its Total: row (=> skip rest, §9.5)

    def close_block():
        if cur is None or cur_is_rollup:
            return
        if not seen_total:
            # no Total: row (§9.2). Valid only if the block had no *activity* — a block may still
            # carry zero-amount transaction rows (e.g. 03-24!1112-01-016-00 has one F=G=0 row).
            sc = sum((t.cargo for t in cur.transactions), ZERO)
            sa = sum((t.abono for t in cur.transactions), ZERO)
            if sc != ZERO or sa != ZERO:
                raise LedgerParseError(
                    f"account {cur.ledger_account} has non-zero activity but no Total: row")
            cur.had_total_row = False
            cur.total_cargos = ZERO
            cur.total_abonos = ZERO
            if cur.transactions and cur.transactions[-1].saldo is not None:
                cur.closing = cur.transactions[-1].saldo
            else:
                cur.closing = cur.opening
        blocks.append(cur)

    r = header_row + 1
    max_row = ws.max_row
    pending_total = None   # summary row (F,G,H) seen just before an empty 'Total:' label row
    while r <= max_row:
        a = _norm(ws.cell(r, cols.fecha).value)
        e = _norm(ws.cell(r, cols.referencia).value)

        # grand total 'Total Bancos :' or spaced 'T o t a l:' (§9.2) => end of data
        e_head = e.split(":")[0].strip()
        is_grand = ("Bancos" in e) or (e_head.startswith("T") and " " in e_head
                                       and e_head.replace(" ", "") == "Total")
        if is_grand:
            break

        if ACCT_RE.match(a):
            close_block()
            acct = a
            cur_is_rollup = acct in rollup
            opening = _cell_num(ws, r, cols.saldo) or ZERO
            cur = LedgerAccountBlock(
                ledger_account=acct,
                label=_norm(ws.cell(r, cols.tipo).value),
                opening=opening, closing=opening,
                total_cargos=ZERO, total_abonos=ZERO,
            )
            seen_total = False
            pending_total = None
            col_i = _cell_num(ws, r, cols.saldo + 1)
            if col_i is not None:
                cur.col_i_annotation.append(col_i)
            r += 1
            continue

        # account total row: col Referencia starts with 'Total' (grand totals broke out above).
        # Some sheets carry the sums on an unlabelled summary row and leave the 'Total:' row blank
        # (e.g. 02-24!1112-01-017-00), so fall back to the pending summary row's values.
        if e.startswith("Total") and cur is not None and not cur_is_rollup and not seen_total:
            row_c = _cell_num(ws, r, cols.cargos)
            row_a = _cell_num(ws, r, cols.abonos)
            row_h = _cell_num(ws, r, cols.saldo)
            if row_c is None and row_a is None and row_h is None and pending_total is not None:
                row_c, row_a, row_h = pending_total
            cur.total_cargos = row_c or ZERO
            cur.total_abonos = row_a or ZERO
            if row_h is not None:
                cur.closing = row_h
            elif cur.transactions and cur.transactions[-1].saldo is not None:
                cur.closing = cur.transactions[-1].saldo
            col_i = _cell_num(ws, r, cols.saldo + 1)
            if col_i is not None:
                cur.col_i_annotation.append(col_i)
            seen_total = True
            r += 1
            continue

        # transaction row
        if DATE_RE.match(a) and cur is not None and not cur_is_rollup and not seen_total:
            cargo = _cell_num(ws, r, cols.cargos) or ZERO
            abono = _cell_num(ws, r, cols.abonos) or ZERO
            # Asset account: Cargos increase the balance, Abonos decrease it (§9.2). §9.2 assumed a
            # row never carries both, but 05-24 row 988 (a transfer with an offset) does; keep both
            # raw values (so the Σcargos/Σabonos tie-out to the Total row holds) and derive the
            # matching direction/amount from the net movement.
            net = cargo - abono
            direction = "inflow" if net >= ZERO else "outflow"
            amount = abs(net)
            txn = LedgerTransaction(
                entity=entity, period=period, ledger_account=cur.ledger_account,
                sheet=sheet_name, row_no=r, txn_date=_es_date(a),
                tipo=_norm(ws.cell(r, cols.tipo).value) or None,
                poliza=_norm(ws.cell(r, cols.numero).value) or None,
                concepto=(ws.cell(r, cols.concepto).value or None),
                referencia=_norm(ws.cell(r, cols.referencia).value) or None,
                cargo=cargo, abono=abono, saldo=_cell_num(ws, r, cols.saldo),
                direction=direction, amount=amount,
            )
            cur.transactions.append(txn)
            r += 1
            continue

        # capture an unlabelled summary row (numeric saldo + cargos/abonos, no date) as the
        # candidate total, in case the following 'Total:' label row is blank
        if cur is not None and not cur_is_rollup and not seen_total and not ACCT_RE.match(a):
            h = _cell_num(ws, r, cols.saldo)
            fc = _cell_num(ws, r, cols.cargos)
            ac = _cell_num(ws, r, cols.abonos)
            if h is not None and (fc is not None or ac is not None):
                pending_total = (fc, ac, h)

        r += 1  # blank / filler / embedded schedule (§9.5) — skipped

    close_block()
    return blocks


class _Grid:
    """In-memory grid mirroring the ws.cell/.max_row/.max_column/.title API used by parse_sheet.

    Built by one sequential pass with read_only iteration (fast); random access on read_only
    worksheets is O(n) per cell and hangs on a 1700-row sheet.
    """

    def __init__(self, ws):
        self.title = ws.title
        self._rows: list[tuple] = [tuple(r) for r in ws.iter_rows(values_only=True)]
        self.max_row = len(self._rows)
        self.max_column = max((len(r) for r in self._rows), default=0)

    class _Cell:
        __slots__ = ("value",)

        def __init__(self, value):
            self.value = value

    def cell(self, r: int, c: int):
        if 1 <= r <= self.max_row:
            row = self._rows[r - 1]
            if 1 <= c <= len(row):
                return self._Cell(row[c - 1])
        return self._Cell(None)


def parse_ledger_file(path: str, entity: str, period: str,
                      cfg: EntityConfig) -> list[LedgerAccountBlock]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = period_to_sheet(period)
    if sheet not in wb.sheetnames:
        raise LedgerParseError(f"sheet {sheet!r} for period {period} not in {wb.sheetnames}")
    grid = _Grid(wb[sheet])
    wb.close()
    return parse_sheet(grid, entity, period, cfg)
